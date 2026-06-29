import aiosqlite
import csv
import hashlib
import io
import re
from datetime import datetime, date, timezone
from app.models.transactions import TransactionCreate, TransactionUpdate


MONTHS = [
    (1, "JAN"), (2, "FEB"), (3, "MAR"), (4, "APR"),
    (5, "MAY"), (6, "JUN"), (7, "JUL"), (8, "AUG"),
    (9, "SEP"), (10, "OCT"), (11, "NOV"), (12, "DEC"),
]

MONTH_MAP = dict(MONTHS)


def date_to_fy_start(year: int, month: int) -> int:
    """Return FY start year. April onward belongs to current year; Jan-Mar belonged to FY started last year."""
    return year if month >= 4 else year - 1


def fy_to_calendar(fy_start: int, month: int) -> int:
    """Compute calendar year from FY start year and month number."""
    return fy_start if month >= 4 else fy_start + 1


def fy_label(fy_start: int) -> str:
    return f"FY {fy_start}-{fy_start + 1}"


def _row_to_dict(row) -> dict | None:
    if row is None:
        return None
    return dict(row)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _parse_amount(val: str) -> float:
    if not val or not val.strip():
        return 0.0
    cleaned = val.strip().replace(",", "").replace("INR", "")
    return float(cleaned)


def _parse_date_str(val: str) -> date:
    val = val.strip()
    for fmt in ("%d/%m/%Y %H:%M:%S.%f", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y",
                "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%d/%m/%y"):
        try:
            return datetime.strptime(val, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse date: {val}")


def compute_txn_hash(txn: dict) -> str:
    parts = [
        str(txn.get("txn_date", "")),
        str(txn.get("narration", "")),
        str(txn.get("reference", "")),
        str(round(float(txn.get("debit", 0) or 0), 2)),
        str(round(float(txn.get("credit", 0) or 0), 2)),
    ]
    raw = "|".join(parts)
    return hashlib.sha256(raw.encode()).hexdigest()


def _find_header_row(all_rows: list[list[str]], max_scan: int = 40) -> int:
    date_headers = {"date", "txn date", "txndate", "transaction date", "txn_date"}
    for i, row in enumerate(all_rows[:max_scan]):
        lower = [c.strip().lower() for c in row]
        if any(h in date_headers for h in lower):
            return i
    return 0


def _detect_format(headers: list[str]) -> dict[str, int]:
    lower = [h.strip().lower() for h in headers]
    mapping = {}

    date_indices = [i for i, h in enumerate(lower) if h in ("date", "txn date", "txndate", "transaction date")]
    value_dt_indices = [i for i, h in enumerate(lower) if "value dt" in h or h in ("value date",)]
    narration_indices = [i for i, h in enumerate(lower) if h in ("narration", "description", "desc", "particulars")]
    ref_indices = [i for i, h in enumerate(lower) if "chq" in h or "ref" in h or h in ("cheque no", "cheque")]

    withdrawal_indices = [i for i, h in enumerate(lower) if "withdrawal" in h]
    deposit_indices = [i for i, h in enumerate(lower) if "deposit" in h or "cr" in h and "dr" not in h]
    balance_indices = [i for i, h in enumerate(lower) if "balance" in h or "closing" in h]

    crdr_indices = [i for i, h in enumerate(lower) if h in ("cr/dr", "transaction type", "type")]
    amount_indices = [i for i, h in enumerate(lower) if h in ("amount (inr)", "amount", "amount(rupees)")]

    category_indices = [i for i, h in enumerate(lower) if h == "category"]

    if date_indices:
        mapping["txn_date"] = date_indices[0]
    if value_dt_indices:
        mapping["value_date"] = value_dt_indices[0]
    if narration_indices:
        mapping["narration"] = narration_indices[0]
    if ref_indices:
        mapping["reference"] = ref_indices[0]
    if category_indices:
        mapping["category"] = category_indices[0]

    if withdrawal_indices and deposit_indices:
        mapping["type"] = "dual"
        mapping["debit"] = withdrawal_indices[0]
        mapping["credit"] = deposit_indices[0]
    elif crdr_indices and amount_indices:
        mapping["type"] = "single"
        mapping["crdr"] = crdr_indices[0]
        mapping["amount"] = amount_indices[0]
    else:
        mapping["type"] = "dual"
        credit_guess = [i for i, h in enumerate(lower) if "credit" in h or "cr" in h]
        debit_guess = [i for i, h in enumerate(lower) if "debit" in h or "dr" in h]
        if credit_guess and debit_guess:
            mapping["credit"] = credit_guess[0]
            mapping["debit"] = debit_guess[0]
        else:
            raise ValueError("Cannot detect transaction format. Headers: " + ", ".join(headers))

    if balance_indices:
        mapping["balance"] = balance_indices[0]

    return mapping


def parse_transactions(content: str, filename: str) -> list[dict]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    rows = []

    if ext == "csv":
        rows = _parse_csv(content)
    elif ext in ("xls", "xlsx"):
        raise ValueError("Binary format — use parse_excel for bytes input")
    else:
        rows = _parse_csv(content)

    return rows


def parse_csv_rows(content: str) -> list[dict]:
    return _parse_csv(content)


def parse_excel_bytes(data: bytes, filename: str) -> list[dict]:
    import openpyxl
    import xlrd

    ext = filename.rsplit(".", 1)[-1].lower()
    raw_rows = []

    if ext == "xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        ws = wb.active
        for ws_row in ws.iter_rows(values_only=True):
            raw_rows.append([str(c) if c is not None else "" for c in ws_row])
    elif ext == "xls":
        wb = xlrd.open_workbook(file_contents=data)
        ws = wb.sheet_by_index(0)
        for r in range(ws.nrows):
            raw_rows.append([str(ws.cell_value(r, c)).strip() for c in range(ws.ncols)])
    else:
        raise ValueError(f"Unsupported Excel format: {ext}")

    if len(raw_rows) < 2:
        return []

    header_idx = _find_header_row(raw_rows)
    headers = raw_rows[header_idx]
    mapping = _detect_format(headers)
    return _map_rows(raw_rows[header_idx + 1:], mapping)


def _parse_csv(content: str) -> list[dict]:
    lines = content.splitlines(True)
    if not lines:
        return []
    reader = csv.reader(lines)
    all_rows = list(reader)
    if len(all_rows) < 2:
        return []

    header_idx = _find_header_row(all_rows)
    headers = all_rows[header_idx]
    mapping = _detect_format(headers)
    return _map_rows(all_rows[header_idx + 1:], mapping)


def _map_rows(raw_rows: list[list[str]], mapping: dict[str, int]) -> list[dict]:
    results = []
    fmt_type = mapping.get("type", "dual")

    for row in raw_rows:
        if not any(cell.strip() for cell in row):
            continue

        txn = _row_at(row, mapping.get("txn_date"))
        val_dt = _row_at(row, mapping.get("value_date"))

        if not txn:
            continue

        try:
            txn_date = _parse_date_str(txn)
        except ValueError:
            continue

        value_date = _parse_date_str(val_dt) if val_dt else txn_date

        narration = _row_at(row, mapping.get("narration"))
        reference = _row_at(row, mapping.get("reference"))

        if fmt_type == "dual":
            debit = _parse_amount(_row_at(row, mapping.get("debit")))
            credit = _parse_amount(_row_at(row, mapping.get("credit")))
        else:
            amount = _parse_amount(_row_at(row, mapping.get("amount")))
            crdr = _row_at(row, mapping.get("crdr")).strip().lower()
            debit = amount if crdr in ("dr", "dr.", "debit", "d") else 0.0
            credit = amount if crdr in ("cr", "cr.", "credit", "c") else 0.0

        balance = _parse_amount(_row_at(row, mapping.get("balance")))

        txn_dict: dict = {
            "txn_date": txn_date.isoformat(),
            "value_date": value_date.isoformat(),
            "narration": narration or "",
            "reference": reference or "",
            "debit": round(debit, 2),
            "credit": round(credit, 2),
            "balance": round(balance, 2),
        }

        if "category" in mapping:
            cat_raw = _row_at(row, mapping["category"])
            if cat_raw:
                txn_dict["category"] = cat_raw

        results.append(txn_dict)

    return results


def _row_at(row: list[str], idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    return row[idx].strip()


async def create_transaction(db: aiosqlite.Connection, data: TransactionCreate) -> dict:
    now = _now()
    cursor = await db.execute(
        """INSERT INTO bank_transactions
           (account_id, txn_date, value_date, narration, reference, debit, credit, balance, category_id, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (data.account_id, data.txn_date.isoformat(), data.value_date.isoformat(),
         data.narration, data.reference, data.debit, data.credit, data.balance,
         data.category_id, now, now),
    )
    await db.commit()
    return await get_transaction(db, cursor.lastrowid)


async def find_existing_txn_hashes(
    db: aiosqlite.Connection, account_id: int, hashes: set[str]
) -> set[str]:
    if not hashes:
        return set()
    placeholders = ",".join(["?"] * len(hashes))
    cursor = await db.execute(
        f"SELECT txn_hash FROM bank_transactions WHERE account_id = ? AND txn_hash IN ({placeholders})",
        (account_id, *hashes),
    )
    return {row["txn_hash"] for row in await cursor.fetchall()}


async def bulk_create_transactions(
    db: aiosqlite.Connection, account_id: int, txns: list[dict], skip_existing: bool = False
) -> dict:
    now = _now()
    rows = []
    
    for t in txns:
        h = compute_txn_hash(t)
        rows.append((account_id, t["txn_date"], t["value_date"],
             t.get("narration", ""), t.get("reference", ""),
             t.get("debit", 0), t.get("credit", 0),
             t.get("balance", 0), t.get("category_id"), now, now, h))
    
    skipped = 0
    if skip_existing:
        hashes = {r[11] for r in rows}
        existing = await find_existing_txn_hashes(db, account_id, hashes)
        rows = [r for r in rows if r[11] not in existing]
        skipped = len(existing)
    
    if rows:
        await db.executemany(
            """INSERT INTO bank_transactions
               (account_id, txn_date, value_date, narration, reference, debit, credit, balance, category_id, created_at, updated_at, txn_hash)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            rows,
        )
    await db.commit()
    return {"count": len(rows), "skipped": skipped}


async def get_transaction(db: aiosqlite.Connection, txn_id: int) -> dict | None:
    cursor = await db.execute("SELECT * FROM bank_transactions WHERE id = ?", (txn_id,))
    return _row_to_dict(await cursor.fetchone())


async def update_transaction(db: aiosqlite.Connection, txn_id: int, data: TransactionUpdate) -> dict | None:
    existing = await get_transaction(db, txn_id)
    if not existing:
        return None
    sets, vals = [], []
    field_map = {
        "txn_date": lambda v: v.isoformat(),
        "value_date": lambda v: v.isoformat(),
        "narration": lambda v: v,
        "reference": lambda v: v,
        "debit": lambda v: v,
        "credit": lambda v: v,
        "balance": lambda v: v,
        "category_id": lambda v: v,
    }
    for key, transform in field_map.items():
        val = getattr(data, key)
        if val is not None:
            sets.append(f"{key} = ?")
            vals.append(transform(val))
    if not sets:
        return existing
    sets.append("updated_at = ?")
    vals.append(_now())
    vals.append(txn_id)
    await db.execute(
        f"UPDATE bank_transactions SET {', '.join(sets)} WHERE id = ?",
        vals,
    )
    await db.commit()
    return await get_transaction(db, txn_id)


async def update_transaction_category(db: aiosqlite.Connection, txn_id: int, category_id: int) -> dict | None:
    existing = await get_transaction(db, txn_id)
    if not existing:
        return None
    cursor = await db.execute(
        "UPDATE bank_transactions SET category_id = ?, updated_at = ? WHERE id = ?",
        (category_id, _now(), txn_id),
    )
    await db.commit()
    return await get_transaction(db, txn_id) if cursor.rowcount > 0 else None


async def delete_transaction(db: aiosqlite.Connection, txn_id: int) -> bool:
    cursor = await db.execute("DELETE FROM bank_transactions WHERE id = ?", (txn_id,))
    await db.commit()
    return cursor.rowcount > 0


async def delete_transactions_by_month(
    db: aiosqlite.Connection, account_id: int, year: int, month: int
) -> int:
    start_date = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end_date = f"{year + 1:04d}-01-01"
    else:
        end_date = f"{year:04d}-{month + 1:02d}-01"
    cursor = await db.execute(
        "DELETE FROM bank_transactions WHERE account_id = ? AND txn_date >= ? AND txn_date < ?",
        (account_id, start_date, end_date),
    )
    await db.commit()
    return cursor.rowcount


async def get_month_summary(
    db: aiosqlite.Connection, account_id: int, year: int, month: int
) -> dict:
    start_date = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end_date = f"{year + 1:04d}-01-01"
    else:
        end_date = f"{year:04d}-{month + 1:02d}-01"
    cursor = await db.execute(
        """SELECT COALESCE(SUM(debit), 0) AS total_debit,
                  COALESCE(SUM(credit), 0) AS total_credit,
                  COUNT(*) AS txn_count
           FROM bank_transactions
           WHERE account_id = ? AND txn_date >= ? AND txn_date < ?""",
        (account_id, start_date, end_date),
    )
    row = await cursor.fetchone()
    return {
        "total_debit": round(row["total_debit"], 2),
        "total_credit": round(row["total_credit"], 2),
        "txn_count": row["txn_count"],
    }


async def list_transactions(
    db: aiosqlite.Connection, account_id: int, year: int, month: int
) -> list[dict]:
    start_date = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end_date = f"{year + 1:04d}-01-01"
    else:
        end_date = f"{year:04d}-{month + 1:02d}-01"
    cursor = await db.execute(
        """SELECT t.*, c.name AS cat_name, c.type AS cat_type
           FROM bank_transactions t
           LEFT JOIN transaction_categories c ON t.category_id = c.id
           WHERE t.account_id = ? AND t.txn_date >= ? AND t.txn_date < ?
           ORDER BY t.txn_date, t.id""",
        (account_id, start_date, end_date),
    )
    return [_row_to_dict(r) for r in await cursor.fetchall()]


async def get_transaction_with_category(db: aiosqlite.Connection, txn_id: int) -> dict | None:
    cursor = await db.execute(
        """SELECT t.*, c.name AS cat_name, c.type AS cat_type
           FROM bank_transactions t
           LEFT JOIN transaction_categories c ON t.category_id = c.id
           WHERE t.id = ?""",
        (txn_id,),
    )
    return _row_to_dict(await cursor.fetchone())


async def list_categories_for_transactions(db: aiosqlite.Connection) -> list[dict]:
    cursor = await db.execute(
        "SELECT id, name, type FROM transaction_categories ORDER BY type, name"
    )
    return [_row_to_dict(r) for r in await cursor.fetchall()]


async def get_available_fy_years(db: aiosqlite.Connection, account_id: int) -> list[int]:
    cursor = await db.execute(
        """SELECT DISTINCT
            CASE WHEN CAST(strftime('%m', txn_date) AS INTEGER) >= 4
                 THEN CAST(strftime('%Y', txn_date) AS INTEGER)
                 ELSE CAST(strftime('%Y', txn_date) AS INTEGER) - 1
            END AS fy_start
           FROM bank_transactions WHERE account_id = ? ORDER BY fy_start""",
        (account_id,),
    )
    return [int(r[0]) for r in await cursor.fetchall()]


async def get_available_months(
    db: aiosqlite.Connection, account_id: int, fy_start: int
) -> list[int]:
    cursor = await db.execute(
        """SELECT DISTINCT CAST(strftime('%m', txn_date) AS INTEGER) AS m
           FROM bank_transactions
           WHERE account_id = ?
             AND (
                 (CAST(strftime('%m', txn_date) AS INTEGER) >= 4
                  AND strftime('%Y', txn_date) = CAST(? AS TEXT))
                 OR
                 (CAST(strftime('%m', txn_date) AS INTEGER) < 4
                  AND strftime('%Y', txn_date) = CAST(? AS TEXT))
             )
            ORDER BY CASE WHEN m >= 4 THEN 0 ELSE 1 END, m""",
        (account_id, fy_start, fy_start + 1),
    )
    return [int(r[0]) for r in await cursor.fetchall()]


async def get_summary(db: aiosqlite.Connection, account_id: int, year: int, month: int) -> dict:
    start_date = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end_date = f"{year + 1:04d}-01-01"
    else:
        end_date = f"{year:04d}-{month + 1:02d}-01"
    cursor = await db.execute(
        """SELECT COALESCE(SUM(debit), 0) AS total_debit,
                  COALESCE(SUM(credit), 0) AS total_credit,
                  COUNT(*) AS txn_count
           FROM bank_transactions
           WHERE account_id = ? AND txn_date >= ? AND txn_date < ?""",
        (account_id, start_date, end_date),
    )
    row = await cursor.fetchone()
    return {
        "total_debit": round(row["total_debit"], 2),
        "total_credit": round(row["total_credit"], 2),
        "txn_count": row["txn_count"],
    }


async def auto_classify_existing_transactions(db: aiosqlite.Connection) -> int:
    cursor = await db.execute("PRAGMA table_info(bank_transactions)")
    cols = [row[1] for row in await cursor.fetchall()]
    if "category_id" not in cols:
        return 0

    cursor = await db.execute(
        "SELECT id, type FROM transaction_categories WHERE name IN (?, ?) AND is_system = 1",
        ("Uncategorized Income", "Uncategorized Expense"),
    )
    rows = await cursor.fetchall()
    cat_map = {}
    for r in rows:
        cat_map[r["type"]] = r["id"]

    updated = 0
    for cat_type, column, op in (
        ("Expense", "debit", ">"),
        ("Income", "credit", ">"),
    ):
        cat_id = cat_map.get(cat_type)
        if not cat_id:
            continue
        cursor = await db.execute(
            f"UPDATE bank_transactions SET category_id = ? WHERE category_id IS NULL AND {column} {op} 0",
            (cat_id,),
        )
        updated += cursor.rowcount
    await db.commit()
    return updated


async def get_uncategorized_category_ids(db: aiosqlite.Connection) -> dict[str, int]:
    cursor = await db.execute(
        "SELECT id, type FROM transaction_categories WHERE name IN (?, ?) AND is_system = 1",
        ("Uncategorized Income", "Uncategorized Expense"),
    )
    result = {}
    for row in await cursor.fetchall():
        result[row["type"]] = row["id"]
    return result


async def get_category_name_map(db: aiosqlite.Connection) -> dict[str, int]:
    cursor = await db.execute(
        "SELECT id, name, type FROM transaction_categories"
    )
    result = {}
    for row in await cursor.fetchall():
        display_name = f"{row['type']}:{row['name']}"
        result[display_name] = row["id"]
    return result


async def apply_rules_to_transactions(db: aiosqlite.Connection, account_id: int, year: int, month: int) -> int:
    from app.services import rules as rule_svc
    start_date = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end_date = f"{year + 1:04d}-01-01"
    else:
        end_date = f"{year:04d}-{month + 1:02d}-01"

    rules = await rule_svc.list_rules(db)
    active_rules = [r for r in rules if r["is_active"]]
    if not active_rules:
        return 0

    cursor = await db.execute(
        """SELECT t.id, t.narration, t.debit, t.credit, t.category_id
           FROM bank_transactions t
           WHERE t.account_id = ? AND t.txn_date >= ? AND t.txn_date < ?""",
        (account_id, start_date, end_date),
    )
    txns = await cursor.fetchall()

    updated = 0
    for txn in txns:
        narration = (txn["narration"] or "").strip()
        if not narration:
            continue
        debit = txn["debit"] or 0
        credit = txn["credit"] or 0

        for rule in active_rules:
            applies_to = rule.get("applies_to", "both")
            if applies_to == "debit" and debit <= 0:
                continue
            if applies_to == "credit" and credit <= 0:
                continue

            match = False
            if rule["match_type"] == "contains":
                match = rule["search_text"].lower() in narration.lower()
            else:
                match = narration.lower() == rule["search_text"].lower()

            if match and txn["category_id"] != rule["category_id"]:
                await db.execute(
                    "UPDATE bank_transactions SET category_id = ?, updated_at = ? WHERE id = ?",
                    (rule["category_id"], _now(), txn["id"]),
                )
                updated += 1
            if match:
                break

    if updated > 0:
        await db.commit()
    return updated
